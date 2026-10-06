"""Pruebas de exportación a Excel, respaldos de la BD y límite de la ficha 360."""
import io
import os
import gzip

from openpyxl import load_workbook


# ======================= Exportación XLSX =======================

def _body():
    return {
        "titulo": "Cartera Activa – Prueba",
        "headers": ["Cuenta", "Cliente", "Saldo", "Mora %", "Desembolso"],
        "keys": ["cuenta", "cliente", "saldo", "mora", "fecha"],
        "rows": [
            {"cuenta": "C-001", "cliente": "=cmd|' /C calc'!A0", "saldo": 1250000,
             "mora": 12.5, "fecha": "2026-03-15"},
            {"cuenta": "C-002", "cliente": "Pérez A***", "saldo": "2.500.000",
             "mora": 0, "fecha": "2026-04-01T10:30:00"},
        ],
        "formatos": {"saldo": "moneda", "mora": "porcentaje", "fecha": "fecha"},
    }


def test_export_xlsx_ok(admin_client, db):
    res = admin_client.post("/api/export/xlsx", json=_body())
    assert res.status_code == 200, res.text
    assert res.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    cd = res.headers["content-disposition"]
    assert "attachment" in cd and "cartera-activa-prueba_" in cd and cd.endswith('.xlsx"')

    wb = load_workbook(io.BytesIO(res.content))
    ws = wb.active
    assert ws.sheet_view.showGridLines is False
    assert ws.freeze_panes == "A5"
    assert ws.auto_filter.ref.startswith("A4:")

    # Fila 1 título, fila 2 subtítulo con el usuario
    assert ws["A1"].value == "Cartera Activa – Prueba"
    assert ws["A1"].font.bold and ws["A1"].font.size == 14
    assert ws["A1"].font.color.rgb.endswith("007A73")
    assert "FIDE Seguros" in ws["A2"].value and "Admin Test" in ws["A2"].value

    # Encabezado teal con letra blanca
    h = ws["A4"]
    assert h.value == "Cuenta"
    assert h.fill.start_color.rgb.endswith("00A79D")
    assert h.font.bold and h.font.color.rgb.endswith("FFFFFF")

    # Datos: celda '=cmd' neutralizada con apóstrofo; moneda sin decimales;
    # porcentaje a fracción; fecha ISO convertida a fecha real.
    assert ws["A5"].value == "C-001"
    assert ws["B5"].value.startswith("'=cmd")
    assert ws["C5"].value == 1250000 and ws["C5"].number_format == "#,##0"
    assert abs(ws["D5"].value - 0.125) < 1e-9 and ws["D5"].number_format == "0.0%"
    assert ws["E5"].number_format == "dd/mm/yyyy" and ws["E5"].value.year == 2026
    assert ws["C6"].value == 2500000  # "2.500.000" colombiano → número
    assert ws.column_dimensions["A"].width >= 10

    # Auditoría
    row = db.execute("SELECT details FROM audit_logs WHERE action='xlsx_export'").fetchone()
    assert row and "filas=2" in row["details"] and "columnas=5" in row["details"]


def test_export_xlsx_413_si_excede_filas(admin_client):
    body = {"titulo": "Grande", "headers": ["a"], "keys": ["a"],
            "rows": [{"a": 1}] * 50_001}
    res = admin_client.post("/api/export/xlsx", json=body)
    assert res.status_code == 413
    assert "máximo" in res.json()["detail"]


def test_export_xlsx_413_si_excede_columnas(admin_client):
    cols = [f"c{i}" for i in range(61)]
    body = {"titulo": "Ancha", "headers": cols, "keys": cols, "rows": []}
    res = admin_client.post("/api/export/xlsx", json=body)
    assert res.status_code == 413


def test_export_xlsx_401_sin_sesion(client):
    res = client.post("/api/export/xlsx", json=_body())
    assert res.status_code == 401


def test_export_xlsx_viewer_puede(viewer_client):
    res = viewer_client.post("/api/export/xlsx", json=_body())
    assert res.status_code == 200


# ======================= Respaldos =======================

def test_run_backup_crea_gz_y_rota(db):
    from app import backup as bk
    import shutil
    d = bk.backup_dir()
    shutil.rmtree(d, ignore_errors=True)

    nombre = bk.run_backup()
    assert bk.BACKUP_NAME_RE.match(nombre)
    ruta = os.path.join(d, nombre)
    assert os.path.isfile(ruta)
    assert not os.path.exists(ruta[:-3]), "el .db intermedio debe borrarse"
    with gzip.open(ruta, "rb") as f:
        assert f.read(16).startswith(b"SQLite format 3")

    # Marca en kv_store → ya no está "due"
    row = db.execute("SELECT value FROM kv_store WHERE key='last_backup_at'").fetchone()
    assert row and row["value"]
    assert bk.backup_due() is False

    # Rotación: tras 16 respaldos solo quedan 14
    for _ in range(16):
        bk.run_backup()
    assert len(bk.list_backups()) == bk.KEEP_LAST == 14
    shutil.rmtree(d, ignore_errors=True)


def test_backup_scheduler_no_corre_en_test():
    from app import backup as bk
    assert os.environ.get("APP_ENV") == "test"
    assert bk.start_backup_scheduler() is None


def test_backup_endpoints_requieren_superadmin(viewer_client):
    assert viewer_client.get("/api/admin/backups").status_code == 403
    assert viewer_client.post("/api/admin/backups/run").status_code == 403
    assert viewer_client.get("/api/admin/backups/fide_20260101_0000.db.gz").status_code == 403


def test_backup_endpoints_superadmin(admin_client, db):
    from app import backup as bk
    import shutil
    shutil.rmtree(bk.backup_dir(), ignore_errors=True)

    res = admin_client.post("/api/admin/backups/run")
    assert res.status_code == 200, res.text
    nombre = res.json()["nombre"]
    assert bk.BACKUP_NAME_RE.match(nombre)

    res = admin_client.get("/api/admin/backups")
    assert res.status_code == 200
    lista = res.json()["backups"]
    assert lista and lista[0]["nombre"] == nombre
    assert {"nombre", "tamano", "fecha"} <= set(lista[0].keys())

    res = admin_client.get(f"/api/admin/backups/{nombre}")
    assert res.status_code == 200
    assert res.headers["content-type"].startswith("application/gzip")
    assert gzip.decompress(res.content)[:16].startswith(b"SQLite format 3")
    row = db.execute("SELECT details FROM audit_logs WHERE action='backup_download'").fetchone()
    assert row and nombre in row["details"]

    # Nombre inexistente pero válido → 404
    assert admin_client.get("/api/admin/backups/fide_19990101_0000.db.gz").status_code == 404
    shutil.rmtree(bk.backup_dir(), ignore_errors=True)


def test_backup_download_rechaza_path_traversal(admin_client):
    for malo in ["..%2F..%2Ffide.db", "..%5Cfide.db", "fide_20260101_0000.db.gz%2F..%2Fx",
                 "audit_fallback.log", "fide.db"]:
        res = admin_client.get(f"/api/admin/backups/{malo}")
        assert res.status_code in (400, 404), (malo, res.status_code)
        # Nunca devuelve un archivo
        assert not res.headers.get("content-type", "").startswith("application/gzip")


# ======================= Límite de la ficha 360 =======================

def test_ficha_limite_30_por_hora_viewer(viewer_client, db):
    for i in range(30):
        res = viewer_client.get("/api/cliente/ficha", params={"identificacion": f"10{i:06d}"})
        assert res.status_code == 200, (i, res.text)
        assert res.json()["encontrado"] is False
    res = viewer_client.get("/api/cliente/ficha", params={"identificacion": "9999999"})
    assert res.status_code == 429
    assert "30 consultas" in res.json()["detail"]
    # Cada consulta quedó auditada (30 ok + el rechazo)
    n = db.execute("SELECT COUNT(*) n FROM audit_logs WHERE action='cliente_ficha'").fetchone()["n"]
    assert n == 30
    n2 = db.execute("SELECT COUNT(*) n FROM audit_logs WHERE action='cliente_ficha_rate_limited'").fetchone()["n"]
    assert n2 == 1


def test_ficha_superadmin_sin_limite(admin_client):
    for i in range(32):
        res = admin_client.get("/api/cliente/ficha", params={"identificacion": f"20{i:06d}"})
        assert res.status_code == 200
