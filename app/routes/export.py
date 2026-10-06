"""Exportación a Excel con formato desde el dashboard.

El frontend ya tiene la tabla que el usuario está viendo (filtrada, ordenada
y con la PII enmascarada según su rol), así que en vez de reconstruir cada
consulta en el servidor recibimos filas + columnas y devolvemos un .xlsx
presentable: título, subtítulo con autor y fecha, encabezados teal, filtros,
paneles congelados y formatos de número colombianos (pesos completos, sin
decimales).

Seguridad:
  - require_auth: cualquier rol puede exportar lo que ya ve en pantalla.
  - Límite de 50.000 filas × 60 columnas (413) para no agotar memoria.
  - Las celdas de texto que empiezan por = + - @ se neutralizan con un
    apóstrofo (inyección de fórmulas CSV/XLSX).
  - Se audita como "xlsx_export" (archivo, filas, columnas).
"""
import io
import re
import logging
import unicodedata
from datetime import datetime, date
from typing import Any

from fastapi import APIRouter, Depends, Request, HTTPException
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.auth.middleware import require_auth
from app.audit import log_audit, get_client_ip

router = APIRouter(prefix="/api/export", tags=["export"])
_log = logging.getLogger("fide.export")

MAX_ROWS = 50_000
MAX_COLS = 60

TEAL_OSCURO = "007A73"
TEAL = "00A79D"
BLANCO = "FFFFFF"
GRIS_BORDE = "D9D9D9"

FORMATOS_VALIDOS = {"moneda", "entero", "porcentaje", "fecha", "texto"}
_XLSX_MEDIA = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_ISO_DATE = re.compile(r"^(\d{4})-(\d{2})-(\d{2})(?:[T ](\d{2}):(\d{2})(?::(\d{2}))?)?")
_FORMULA_PREFIX = ("=", "+", "-", "@")


class ExportBody(BaseModel):
    titulo: str = Field(..., min_length=1, max_length=120)
    hoja: str | None = Field(default=None, max_length=31)
    headers: list[str]
    keys: list[str]
    rows: list[dict[str, Any]]
    formatos: dict[str, str] | None = None


def _slug(texto: str) -> str:
    """Nombre de archivo seguro: sin tildes, minúsculas, guiones."""
    nfkd = unicodedata.normalize("NFKD", texto)
    plano = "".join(c for c in nfkd if not unicodedata.combining(c))
    plano = re.sub(r"[^A-Za-z0-9]+", "-", plano).strip("-").lower()
    return plano[:60] or "exportacion"


def _nombre_hoja(texto: str) -> str:
    """Excel prohíbe []:*?/\\ y máximo 31 caracteres en el nombre de hoja."""
    limpio = re.sub(r"[\[\]:*?/\\]", " ", texto).strip()
    return (limpio or "Datos")[:31]


def _parse_fecha(valor):
    """Convierte 'yyyy-mm-dd' o ISO con hora a date/datetime; si no, None."""
    if isinstance(valor, (datetime, date)):
        return valor
    if not isinstance(valor, str):
        return None
    m = _ISO_DATE.match(valor.strip())
    if not m:
        return None
    try:
        y, mo, d = int(m.group(1)), int(m.group(2)), int(m.group(3))
        if m.group(4) is not None:
            return datetime(y, mo, d, int(m.group(4)), int(m.group(5)), int(m.group(6) or 0))
        return date(y, mo, d)
    except ValueError:
        return None


def _a_numero(valor):
    """Intenta convertir a número; respeta None y devuelve None si no se puede."""
    if valor is None or valor == "":
        return None
    if isinstance(valor, bool):
        return int(valor)
    if isinstance(valor, (int, float)):
        return valor
    if isinstance(valor, str):
        s = valor.strip().replace("$", "").replace(" ", "")
        # Formato colombiano "1.234.567,89" → 1234567.89
        if re.fullmatch(r"-?\d{1,3}(\.\d{3})+(,\d+)?", s):
            s = s.replace(".", "").replace(",", ".")
        elif re.fullmatch(r"-?\d+(,\d+)?", s):
            s = s.replace(",", ".")
        s = s.rstrip("%")
        try:
            f = float(s)
            return int(f) if f.is_integer() else f
        except ValueError:
            return None
    return None


def _neutralizar(texto: str) -> str:
    """Anteponer apóstrofo a textos que Excel interpretaría como fórmula."""
    if texto and texto[0] in _FORMULA_PREFIX:
        return "'" + texto
    return texto


def _celda_valor(valor, formato: str):
    """Devuelve (valor_celda, number_format) según el formato pedido."""
    if formato in ("moneda", "entero"):
        n = _a_numero(valor)
        if n is not None:
            return round(n), "#,##0"
    elif formato == "porcentaje":
        n = _a_numero(valor)
        if n is not None:
            # Si viene como 12.5 (por ciento) lo pasamos a fracción; si ya es
            # fracción (0.125) lo dejamos.
            if isinstance(valor, str) and valor.strip().endswith("%"):
                n = n / 100
            elif abs(n) > 1:
                n = n / 100
            return n, "0.0%"
    elif formato == "fecha":
        f = _parse_fecha(valor)
        if f is not None:
            return f, "dd/mm/yyyy"
    elif formato == "texto":
        if valor is None:
            return None, "@"
        return _neutralizar(str(valor)), "@"

    # Sin formato explícito: inferir
    if valor is None:
        return None, "General"
    if isinstance(valor, bool):
        return "Sí" if valor else "No", "General"
    if isinstance(valor, (int, float)):
        return valor, "General"
    if isinstance(valor, (datetime, date)):
        return valor, "dd/mm/yyyy"
    if isinstance(valor, str):
        f = _parse_fecha(valor)
        if f is not None and len(valor.strip()) >= 10:
            return f, "dd/mm/yyyy"
        return _neutralizar(valor), "General"
    # dict/list u otros: representarlos como texto
    return _neutralizar(str(valor)), "General"


def _ancho_texto(valor) -> int:
    if valor is None:
        return 0
    if isinstance(valor, (datetime, date)):
        return 10
    if isinstance(valor, float):
        return len(f"{valor:,.0f}")
    if isinstance(valor, int):
        return len(f"{valor:,}")
    return len(str(valor))


def build_workbook(body: ExportBody, usuario: str) -> Workbook:
    """Arma el libro con el formato institucional. Separado del endpoint
    para poder probarlo sin HTTP."""
    wb = Workbook()
    ws = wb.active
    ws.title = _nombre_hoja(body.hoja or body.titulo)
    ws.sheet_view.showGridLines = False

    formatos = {k: v for k, v in (body.formatos or {}).items() if v in FORMATOS_VALIDOS}
    n_cols = len(body.headers)
    ultima_col = get_column_letter(max(n_cols, 1))

    fuente_base = Font(name="Calibri", size=11)
    borde_fino = Side(style="thin", color=GRIS_BORDE)
    borde = Border(left=borde_fino, right=borde_fino, top=borde_fino, bottom=borde_fino)

    # Fila 1: título
    ws["A1"] = _neutralizar(body.titulo)
    ws["A1"].font = Font(name="Calibri", size=14, bold=True, color=TEAL_OSCURO)
    if n_cols > 1:
        ws.merge_cells(f"A1:{ultima_col}1")
    ws.row_dimensions[1].height = 22

    # Fila 2: subtítulo
    ahora = datetime.now().strftime("%d/%m/%Y %H:%M")
    ws["A2"] = f"FIDE Seguros · generado {ahora} por {usuario}"
    ws["A2"].font = Font(name="Calibri", size=10, italic=True, color="595959")
    if n_cols > 1:
        ws.merge_cells(f"A2:{ultima_col}2")

    # Fila 4: encabezados
    fill_teal = PatternFill("solid", start_color=TEAL, end_color=TEAL)
    fuente_header = Font(name="Calibri", size=11, bold=True, color=BLANCO)
    anchos = []
    for i, h in enumerate(body.headers, start=1):
        c = ws.cell(row=4, column=i, value=_neutralizar(str(h)))
        c.font = fuente_header
        c.fill = fill_teal
        c.border = borde
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        anchos.append(len(str(h)))
    ws.row_dimensions[4].height = 20

    # Datos desde la fila 5
    fila = 5
    for r in body.rows:
        for i, k in enumerate(body.keys, start=1):
            valor, numfmt = _celda_valor(r.get(k), formatos.get(k, ""))
            c = ws.cell(row=fila, column=i, value=valor)
            c.font = fuente_base
            c.border = borde
            c.number_format = numfmt
            if isinstance(valor, (int, float)) and not isinstance(valor, bool):
                c.alignment = Alignment(horizontal="right")
            elif isinstance(valor, (datetime, date)):
                c.alignment = Alignment(horizontal="center")
            anchos[i - 1] = max(anchos[i - 1], _ancho_texto(valor))
        fila += 1

    # Anchos de columna según contenido (mín 10, máx 45)
    for i, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(10, min(45, ancho + 2))

    # Filtros y paneles congelados bajo el encabezado
    if n_cols:
        ultima_fila = max(fila - 1, 4)
        ws.auto_filter.ref = f"A4:{ultima_col}{ultima_fila}"
    ws.freeze_panes = "A5"
    return wb


@router.post("/xlsx")
def export_xlsx(body: ExportBody, request: Request, user=Depends(require_auth)):
    """Genera un .xlsx con formato a partir de la tabla que el usuario ve.

    Los valores llegan ya enmascarados desde el frontend según rol; aquí no
    se consulta la BD ni se vuelve a enmascarar."""
    if len(body.headers) != len(body.keys):
        raise HTTPException(status_code=400,
                            detail="headers y keys deben tener la misma longitud.")
    if not body.headers:
        raise HTTPException(status_code=400, detail="Se requiere al menos una columna.")
    if len(body.rows) > MAX_ROWS or len(body.headers) > MAX_COLS:
        raise HTTPException(
            status_code=413,
            detail=f"La exportación supera el máximo permitido "
                   f"({MAX_ROWS:,} filas × {MAX_COLS} columnas).".replace(",", "."))

    usuario = user.get("display_name") or user.get("username") or "usuario"
    wb = build_workbook(body, usuario)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    nombre = f"{_slug(body.titulo)}_{date.today().isoformat()}.xlsx"
    ip = get_client_ip(request) or "unknown"
    log_audit(user["user_id"], user["username"], "xlsx_export",
              # filas con separador de miles: el scrubber de auditoría hashea cifras de
              # 5+ dígitos seguidos y "50000" quedaría como id=hash.
              f"archivo={nombre} filas={len(body.rows):,}".replace(",", ".")
              + f" columnas={len(body.headers)}", ip)

    return StreamingResponse(
        buf,
        media_type=_XLSX_MEDIA,
        headers={"Content-Disposition": f'attachment; filename="{nombre}"'},
    )
